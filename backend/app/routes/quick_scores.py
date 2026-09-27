"""Quick-analysis (Yahoo-only) scores — same filters/sorting as
/api/company-scores but over `fa_quick_scores`, a separate table: quick and
full-pipeline scores are never mixed. See app/quick_analysis/README.md for
how closely they track (overall error ~1.6 points on 18 analysed companies).

    GET /api/quick-scores?min_growth=70&max_valuation=40&sector=Healthcare&sort_by=overall
    GET /api/quick-scores/export.xlsx      (every scored stock, one sheet)
"""
from __future__ import annotations

from datetime import date
from io import BytesIO

from fastapi import APIRouter, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.infrastructure.database.client import get_db
from app.infrastructure.database.models import QuickScore, Stock
from app.routes.company_scores import _row_dict, list_scores

router = APIRouter(prefix="/api/quick-scores")

_COLUMNS = [
    ("Symbol", "symbol"), ("Company", "company_name"), ("Macro sector", "macro_sector"), ("Sector", "sector"),
    ("Sector framework", "sector_framework"), ("Market cap (Cr)", "market_cap_cr"),
    ("Overall", "overall"), ("Growth", "growth"), ("Growth (Annual)", "growth_annual"),
    ("Growth (Quarterly)", "growth_quarterly"), ("Profitability", "profitability"),
    ("Cash flow", "cash_flow"), ("Balance sheet", "balance_sheet"), ("Efficiency", "efficiency"),
    ("Valuation", "valuation"), ("Rating", "overall_rating"), ("Valuation view", "valuation_view"),
    ("Red flags", "red_flags_text"), ("Latest FY", "latest_fy"), ("Scored on", "scored_on"),
]

_EXTRA_NUMERIC_FIELDS = frozenset({"growth_annual", "growth_quarterly"})


@router.get("")
def list_quick_scores(request: Request):
    return list_scores(request, QuickScore, extra_numeric_fields=_EXTRA_NUMERIC_FIELDS)


@router.get("/export.xlsx")
def export_quick_scores():
    db = get_db()
    try:
        rows = (
            db.query(QuickScore, Stock).join(Stock, Stock.id == QuickScore.stock_id)
            .filter(Stock.is_active.is_(True))
            .order_by(QuickScore.overall.desc().nulls_last(), Stock.symbol).all()
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "Quick scores"
        ws.append([name for name, _ in _COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for score, stock in rows:
            d = _row_dict(score, stock)
            d["market_cap_cr"] = round(d["market_cap"] / 1e7, 1) if d["market_cap"] is not None else None
            d["red_flags_text"] = "; ".join(str(f.get("flag") or f.get("name") or f) if isinstance(f, dict) else str(f)
                                            for f in d["red_flags"])
            d["scored_on"] = (d["scored_at"] or "")[:10]
            ws.append([d.get(key) for _, key in _COLUMNS])
        for i, (name, key) in enumerate(_COLUMNS, 1):
            width = 34 if key == "company_name" else 40 if key == "red_flags_text" else max(len(name) + 2, 12)
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = f"quick_scores_{date.today().isoformat()}.xlsx"
        return StreamingResponse(
            buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'},
        )
    finally:
        db.close()

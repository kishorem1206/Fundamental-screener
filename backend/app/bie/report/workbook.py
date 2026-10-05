"""The report's model as an Excel workbook: history, assumptions, the
forecast statements, the valuation and the exhibits as native charts.

The cells hold values, not formulas: the model is calculated by the engine
(so that analyst overrides, sources and checks stay in one place) and this
file is its print-out. To change an assumption, use the Assumptions page and
download again.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.bie.report.builder import build_report

CR = 1e7
_HEAD = PatternFill("solid", fgColor="16263D")
_TOTAL = PatternFill("solid", fgColor="EEF1F4")
_OUT = Path(__file__).resolve().parents[3] / "reports"


class _Sheet:
    def __init__(self, book: Workbook, title: str, subtitle: str = ""):
        self.ws = book.create_sheet(title[:31])
        self.ws["A1"] = title.replace("_", " ")
        self.ws["A1"].font = Font(bold=True, size=13)
        self.ws["A2"] = subtitle
        self.ws["A2"].font = Font(italic=True, color="666666")
        self.row = 4
        self.ws.column_dimensions["A"].width = 46

    def header(self, *cells) -> None:
        for c, value in enumerate(cells, 1):
            cell = self.ws.cell(self.row, c, value)
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), _HEAD
            cell.alignment = Alignment(horizontal="left" if c == 1 else "right", wrap_text=True)
            if c > 1:
                self.ws.column_dimensions[get_column_letter(c)].width = max(self.ws.column_dimensions[get_column_letter(c)].width or 0, 15)
        self.row += 1

    def line(self, label, *values, total: bool = False, fmt: str = "#,##0") -> None:
        self.ws.cell(self.row, 1, label).font = Font(bold=total)
        for c, value in enumerate(values, 2):
            cell = self.ws.cell(self.row, c, value)
            if isinstance(value, (int, float)):
                cell.number_format = fmt
            cell.font = Font(bold=total)
            if total:
                cell.fill = _TOTAL
        if total:
            self.ws.cell(self.row, 1).fill = _TOTAL
        self.row += 1

    def gap(self, text: str | None = None) -> None:
        self.row += 1
        if text:
            self.ws.cell(self.row, 1, text).font = Font(bold=True, color="1F4E79")
            self.row += 1


def _c(x):
    return None if x is None else x / CR


def build_workbook(db: Session, company_id: str) -> str:
    r = build_report(db, company_id)
    db.commit()
    stock, v = r["stock"], r["valuation"]
    book = Workbook()
    book.remove(book.active)
    modelled = bool(v) and "skipped" not in v
    base = v["scenarios"]["Base"] if modelled else None
    rows = base["rows"] if modelled else []
    years = [y["label"] for y in rows]
    annual = sorted(r["annual"], key=lambda a: a["end"])

    s = _Sheet(book, "Read Me", f"{stock.company_name} ({stock.symbol}) · generated {r['generated']:%d %b %Y} · ₹ crore unless noted")
    for text in ("This workbook is the print-out of the model behind the company intelligence report.",
                 "Cells hold values, not formulas: the engine calculates the model so that sources, overrides and checks stay in one place.",
                 "To change an assumption, use the Assumptions page of the app (every override needs a reason) and download again.",
                 "History comes from exchange filings listed on the Source Register sheet. Forecasts are rule-based estimates, not company guidance.",
                 "Nothing here is a recommendation to buy, sell or hold."):
        s.line(text)
    if modelled:
        s.gap("At a glance")
        s.header("Measure", "Value")
        s.line("Basis of the model", v["basis"].title())
        s.line("Last full year", f"FY{v['fy'].year}")
        s.line("Share price used (₹)", v["price"], fmt="#,##0.00")
        for key, label in (("dcf", "Discounted cash flow (₹ per share)"), ("sotp", "Sum of the parts (₹ per share)"),
                           ("relative", "Peer multiple (₹ per share)"), ("average", "Blended reference (₹ per share)")):
            if key in base:
                s.line(label, base[key], fmt="#,##0.00", total=key == "average")
        s.line("Analyst overrides in force", len(v["overrides"]))
        worst = max((abs(y["sheet"]["residual"]) for y in rows if "sheet" in y), default=None)
        if worst is not None:
            s.line("Largest balance-sheet difference in the forecast", worst / CR, fmt="0.00")

    s = _Sheet(book, "Source Register", "Every document a figure was read from")
    s.header("No.", "Title", "Publisher", "Date", "Link")
    for src in r["sources"]:
        s.line(src["n"], src["title"], src["publisher"], str(src["date"] or ""), src["readable_url"] or src["url"])
    s.ws.column_dimensions["A"].width, s.ws.column_dimensions["B"].width, s.ws.column_dimensions["E"].width = 6, 70, 90

    s = _Sheet(book, "Historic_IS", f"Reported results, {r['basis'] or ''}")
    s.header("₹ crore", *[a["label"] for a in annual])
    for key, label, fmt in (("revenue", "Revenue from operations", "#,##0"), ("excise", "Excise duty", "#,##0"), ("net_revenue", "Revenue net of excise", "#,##0"),
                            ("other_income", "Other income", "#,##0"), ("pbt_before_exceptional", "Profit before exceptional items and tax", "#,##0"),
                            ("exceptional", "Exceptional items", "#,##0"), ("pbt", "Profit before tax", "#,##0"), ("tax", "Tax expense", "#,##0"),
                            ("profit_continuing", "Profit from continuing operations", "#,##0"), ("profit_discontinued", "Profit from discontinued operations", "#,##0"),
                            ("profit", "Profit after tax", "#,##0"), ("eps", "Earnings per share, basic (₹)", "#,##0.00")):
        if any(a.get(key) is not None for a in annual):
            s.line(label, *[a.get(key) for a in annual], fmt=fmt, total=key in ("net_revenue", "profit"))
    quarters = sorted(r.get("quarters") or [], key=lambda q: q["end"])
    if quarters:
        s.gap("Latest quarters")
        s.header("₹ crore", *[q["label"] for q in quarters])
        for key, label in (("revenue", "Revenue"), ("net_revenue", "Revenue net of excise"), ("pbt", "Profit before tax"), ("profit", "Profit after tax")):
            if any(q.get(key) is not None for q in quarters):
                s.line(label, *[q.get(key) for q in quarters])

    if r["segments"]:
        s = _Sheet(book, "Historic_Segments", "Segments as reported, before inter-segment elimination")
        for table in r["segments"]:
            s.gap(f"{table['label']}, {table['basis'].title()}")
            s.header("Segment", "Revenue", "Result", "Margin", "Share of revenue", "Share of result", "Capital employed", "Return on capital")
            for row in table["rows"]:
                s.line(row["name"], row["revenue"], row["result"], row["margin"], row["revenue_share"], row["result_share"], row["capital_employed"], row["return_on_capital"])
        for c in "DEFH":
            for cell in s.ws[c][4:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0%"

    cash = r["cash"]["rows"] if r.get("cash") else []
    if cash:
        s = _Sheet(book, "Historic_CF_BS", "Cash flow and balance-sheet markers from each year's own filing")
        s.header("₹ crore", *[c["label"] for c in cash])
        for key, label in (("cfo", "Operating cash flow"), ("capex", "Capital expenditure"), ("fcf", "Free cash flow"), ("dividends", "Dividends paid"),
                           ("assets", "Total assets"), ("equity", "Total equity"), ("inventories", "Inventories"), ("receivables", "Trade receivables"),
                           ("current_investments", "Current investments"), ("net_liquid", "Cash and current investments less borrowings")):
            s.line(label, *[c.get(key) for c in cash], total=key == "fcf")

    if modelled:
        s = _Sheet(book, "Assumptions", "The model's own estimate for every input that is a choice, with the rule that produced it")
        s.header("Assumption", "Applies to", "Value", "Unit", "Rule")
        for a in v["assumptions"]:
            s.line(a["key"], a["for"], a["value"], a["unit"], a["why"], fmt="0.0000")
        s.ws.column_dimensions["B"].width, s.ws.column_dimensions["E"].width = 34, 120
        if v["overrides"]:
            s.gap("Analyst overrides in force")
            s.header("Assumption", "Applies to", "Model", "Override", "Reason")
            for o in v["overrides"]:
                s.line(o["metric"], f"{o['unit']} {o.get('year') or ''}".strip(), o["system"], o["value"], o["reason"], fmt="0.0000")

        s = _Sheet(book, "Scenarios", "Bull and Bear move each unit by how much its own record has varied")
        s.header("Unit", "Growth shift (points)", "Margin shift (proportion)", "Year-1 growth, Base", "Margin, Base")
        for u in v["units"]:
            s.line(u["name"], u["g_shift"], u["m_shift"], u["start"], u["margin"], fmt="0.0%")
        s.gap("Profit after tax by scenario")
        s.header("₹ crore", *years)
        for name in ("Bear", "Base", "Bull"):
            s.line(name, *[_c(y["pat"]) for y in v["scenarios"][name]["rows"]], total=name == "Base")

        s = _Sheet(book, "Segments", "Base-case forecast by unit")
        s.header("₹ crore", *years)
        for u in rows[0]["units"]:
            s.line(f"{u['name']}: revenue", *[_c(next(x for x in y["units"] if x["name"] == u["name"])["revenue"]) for y in rows])
            s.line(f"{u['name']}: result", *[_c(next(x for x in y["units"] if x["name"] == u["name"])["result"]) for y in rows])

        s = _Sheet(book, "Income_Statement", "Base case")
        s.header("₹ crore", *years)
        s.line("Revenue", *[_c(y["revenue"]) for y in rows], total=True)
        if v["has_excise"]:
            s.line("Excise duty", *[_c(y["excise"]) for y in rows])
            s.line("Revenue net of excise", *[_c(y["net_revenue"]) for y in rows], total=True)
        s.line("Profit before tax", *[_c(y["pbt"]) for y in rows])
        s.line("Profit after tax", *[_c(y["pat"]) for y in rows], total=True)
        s.line("Share of associates' profit", *[_c(y["associates"]) for y in rows])
        s.line("Less minority shareholders' share", *[_c(-y["minority_profit"]) for y in rows])
        s.line("Profit belonging to shareholders", *[_c(y["owners"]) for y in rows], total=True)
        s.line("Earnings per share (₹)", *[y["eps"] for y in rows], fmt="#,##0.00")
        if rows[0].get("dps") is not None:
            s.line("Dividend per share (₹)", *[y["dps"] for y in rows], fmt="#,##0.00")

        if rows[0].get("cfo") is not None:
            s = _Sheet(book, "Working_Capital", "Each line held at its share of revenue in the last full year")
            s.header("₹ crore", *years)
            for line in rows[0]["working_capital"]:
                s.line(("Less " + line["label"].lower()) if line["sign"] < 0 else line["label"],
                       *[_c(next(x for x in y["working_capital"] if x["label"] == line["label"])["value"] * line["sign"]) for y in rows])
            s.line("Operating working capital", *[_c(y["nwc"]) for y in rows], total=True)
            s.line("Change in the year", *[_c(y["nwc_change"]) for y in rows])

            s = _Sheet(book, "Cash_Flow", "From profit to cash, and where the cash goes")
            s.header("₹ crore", *years)
            s.line("Profit after tax", *[_c(y["pat"]) for y in rows])
            s.line("Add depreciation and amortisation", *[_c(y["da"]) for y in rows])
            s.line("Less increase in working capital", *[_c(-y["nwc_change"]) for y in rows])
            s.line("Operating cash flow", *[_c(y["cfo"]) for y in rows], total=True)
            s.line("Less capital expenditure", *[_c(-y["capex"]) for y in rows])
            s.line("Less dividends", *[_c(-y["dividends"]) for y in rows])
            s.line("Less acquisitions completed since the balance sheet", *[_c(-y["acquisitions"]) for y in rows])
            s.line("Cash left", *[_c(y["cash_after"]) for y in rows], total=True)
            s.line("Cash and current investments at year end", *[_c(y["liquid_close"]) for y in rows])

        if v.get("sheet0"):
            s0 = v["sheet0"]
            s = _Sheet(book, "Balance_Sheet", "Grouped the way the forecast moves it")
            s.header("₹ crore, year end", f"FY{v['fy'].year % 100:02d}", *years)
            for key, label, total in (("liquid", "Cash and current investments", False), ("wc_assets", "Receivables, inventories, other current assets", False),
                                      ("fixed", "Fixed and other long-term assets", False), ("investments", "Long-term investments and associates", False),
                                      ("assets", "Total assets", True), ("wc_liabilities", "Payables and other current liabilities", False),
                                      ("borrowings", "Borrowings", False), ("other_liabilities", "Other liabilities", False), ("equity", "Equity", False)):
                s.line(label, _c(s0[key]), *[_c(y["sheet"][key]) for y in rows], total=total)
            s.line("Total liabilities and equity", _c(s0["wc_liabilities"] + s0["borrowings"] + s0["other_liabilities"] + s0["equity"]),
                   *[_c(y["sheet"]["assets"] - y["sheet"]["residual"]) for y in rows], total=True)
            s.line("Difference (must be zero)", 0.0, *[_c(y["sheet"]["residual"]) for y in rows], fmt="0.00")

            s = _Sheet(book, "Model_Checks", "Each check is the largest difference over the forecast years; zero is a pass")
            s.header("Check", "Largest difference, ₹ crore", "Result")
            checks = (("Balance sheet: assets equal liabilities plus equity", max(abs(y["sheet"]["residual"]) for y in rows)),
                      ("Operating cash flow equals profit plus depreciation less working-capital change",
                       max(abs(y["cfo"] - (y["pat"] + y["da"] - y["nwc_change"])) for y in rows)),
                      ("Working-capital lines add up to the total", max(abs(sum(x["value"] * x["sign"] for x in y["working_capital"]) - y["nwc"]) for y in rows)),
                      ("Year-end cash rolls forward from cash left", max(abs(b["liquid_close"] - a["liquid_close"] - b["cash_after"]) for a, b in zip(rows, rows[1:]))))
            for label, worst in checks:
                s.line(label, worst / CR, "Pass" if worst / CR < 0.5 else "FAIL", fmt="0.00")
            s.line("Cash stays positive", None, "Pass" if all(y["liquid_close"] >= 0 for y in rows) else "Warning")
            s.line("Dividend covered by cash after capex", None, "Pass" if not any(y["dividend_uncovered"] for y in rows) else "Warning")

        c = v.get("cost_of_capital")
        if c and "dcf" in base:
            s = _Sheet(book, "Cost_Of_Capital", "Capital asset pricing model")
            s.header("Input", "Value")
            for label, key in (("Ten-year government yield", "risk_free"), ("Equity risk premium, India", "equity_risk_premium"), ("Beta", "beta"),
                               ("Cost of equity", "cost_of_equity"), ("Debt share of capital", "debt_weight"), ("Pre-tax cost of debt", "cost_of_debt"),
                               ("Cost of capital", "wacc"), ("Terminal growth", "terminal_growth")):
                s.line(label, c[key], fmt="0.00" if key == "beta" else "0.00%", total=key == "wacc")
            s.line("Terminal return on capital", v["roc"], fmt="0.0%")

            s = _Sheet(book, "DCF", "Free cash flow to the firm, Base case")
            s.header("₹ crore", *years)
            s.line("Operating profit after tax", *[_c(y["nopat"]) for y in rows])
            s.line("Add depreciation and amortisation", *[_c(y["da"]) for y in rows])
            s.line("Less capital expenditure", *[_c(-y["capex"]) for y in rows])
            s.line("Less increase in working capital", *[_c(-y["nwc_change"]) for y in rows])
            s.line("Free cash flow to the firm", *[_c(y["fcff"]) for y in rows], total=True)
            s.line("Present value", *[_c(x) for x in base["dcf_pv"]])
            s.gap("From the business to the shares")
            s.header("Item", "₹ crore", "Basis")
            s.line("Value of the business (forecast years + terminal value)", _c(base["dcf_ev"]), f"{base['dcf_terminal_share']:.0%} from the terminal year", total=True)
            s.line("Less minority shareholders' share", _c(-base["dcf_ev"] * v["equity_bridge"]["minority"]), f"{v['equity_bridge']['minority']:.1%}")
            for item in v["equity_bridge"]["items"]:
                s.line(item["label"], _c(item["value"]), item["basis"])
            s.line("Shares outstanding (crore)", _c(v["shares"]), fmt="#,##0.00")
            s.line("Value per share (₹)", base["dcf"], fmt="#,##0.00", total=True)
            s.ws.column_dimensions["C"].width = 60

        if base.get("sotp_parts"):
            s = _Sheet(book, "SOTP", "Next year's segment result × the multiple of its line of business")
            s.header("Segment", f"{years[0]} result", "EV ÷ EBIT", "Value", "Where the multiple comes from")
            for part in base["sotp_parts"]:
                s.line(part["name"], _c(part["result"]), part["multiple"], _c(part["value"]), part["basis"])
            s.line("Less unallocated costs", None, None, _c(-base["sotp_unallocated"]))
            s.line("Assets outside the forecast, net", None, None, _c(v["equity_bridge"]["added"]))
            s.line("Value per share (₹)", None, None, base["sotp"], total=True)
            s.ws.column_dimensions["E"].width = 80

        if "relative" in base:
            s = _Sheet(book, "Relative_Valuation", "Peers' price ÷ earnings applied to next year's earnings per share")
            if v.get("pe_parts"):
                s.header("Segment", "Share of profit", "Peer multiple", "Contribution", "Basis")
                for part in v["pe_parts"]:
                    s.line(part["name"], part["weight"], part["multiple"], part["weight"] * part["multiple"], part["basis"], fmt="0.00")
            s.gap()
            s.header("Measure", "Value")
            s.line("Multiple applied (×)", v["peer_pe"], fmt="0.00")
            s.line(f"Earnings per share, {years[0]} (₹)", rows[0]["eps"], fmt="#,##0.00")
            s.line("Value per share (₹)", base["relative"], fmt="#,##0.00", total=True)

        s = _Sheet(book, "Sensitivities", "Base case, ₹ per share")
        if v.get("sensitivity"):
            s.gap("Cash-flow value: terminal growth down, cost of capital across")
            s.header("", *[f"{w:.2%}" for w in v["sensitivity"]["waccs"]])
            for row in v["sensitivity"]["rows"]:
                s.line(f"{row['g']:.1%}", *row["values"])
        if v.get("sotp_sensitivity"):
            ss = v["sotp_sensitivity"]
            s.gap(f"Sum of the parts: {ss['down']} multiple down, {ss['across']} multiple across")
            s.header("", *[f"{m:.1f}×" for m in ss["across_multiples"]])
            for row in ss["rows"]:
                s.line(f"{row['multiple']:.1f}×", *row["values"])

        s = _Sheet(book, "Valuation_Summary", f"Against ₹{v['price']:,.2f} on {v['price_date']:%d %b %Y}")
        s.header("₹ per share", "Discounted cash flow", "Sum of the parts", "Peer multiple", "Blended reference", "Blend against price")
        for name in ("Bear", "Base", "Bull"):
            sc = v["scenarios"][name]
            s.line(name, sc.get("dcf"), sc.get("sotp"), sc.get("relative"), sc.get("average"), sc.get("gap"), total=name == "Base")
        for cell in s.ws["F"][4:]:
            if isinstance(cell.value, float):
                cell.number_format = "+0.0%;-0.0%"
        s.gap("Blend weights")
        for key, label in (("dcf", "Discounted cash flow"), ("sotp", "Sum of the parts"), ("relative", "Peer multiple")):
            if key in base.get("weights", {}):
                s.line(label, base["weights"][key], fmt="0%")

    s = _Sheet(book, "Peers", "Latest full year from each company's own filing")
    s.header("Company", "Peer for segment", "Year", "Market value", "Revenue", "Profit after tax", "Net margin")
    for p in r["peers"]:
        s.line(p["name"] + (" (subject)" if p["is_subject"] else ""), p["for_segment"] or "", p["year"], p["market_cap"], p["revenue"], p["profit"], p["margin"])
    for cell in s.ws["G"][4:]:
        if isinstance(cell.value, float):
            cell.number_format = "0.0%"

    kpis = r.get("sector_measures") or []
    if kpis:
        s = _Sheet(book, "Sector_Measures", "As stated by the company in its own filings")
        s.header("Measure", "Value", "Unit", "As stated", "Filed", "Document", "Page")
        for k in kpis:
            s.line(k["label"], k["value"], k["unit"], k["quote"], str(k["date"]), k["document"], k["page"], fmt="#,##0.00")
        s.ws.column_dimensions["D"].width, s.ws.column_dimensions["F"].width = 70, 50

    if r.get("exhibits"):
        data = _Sheet(book, "Chart_Data", "One block per exhibit")
        charts = book.create_sheet("Charts")
        charts["A1"] = "Exhibits"
        charts["A1"].font = Font(bold=True, size=13)
        for index, e in enumerate(r["exhibits"]):
            data.gap(f"Exhibit {e['n']}. {e['title']} ({e['unit']})")
            data.header("", *e["labels"])
            first = data.row
            for series in e["series"]:
                data.line(series["name"], *series["values"], fmt="#,##0.00")
            chart = BarChart() if e["kind"] == "bar" else LineChart()
            chart.title, chart.height, chart.width = f"Exhibit {e['n']}. {e['title']}", 7.5, 15
            chart.y_axis.title = e["unit"]
            chart.add_data(Reference(data.ws, min_col=1, max_col=len(e["labels"]) + 1, min_row=first, max_row=data.row - 1), from_rows=True, titles_from_data=True)
            chart.set_categories(Reference(data.ws, min_col=2, max_col=len(e["labels"]) + 1, min_row=first - 1))
            charts.add_chart(chart, f"{'A' if index % 2 == 0 else 'K'}{3 + (index // 2) * 16}")

    _OUT.mkdir(exist_ok=True)
    path = _OUT / f"BIE-{stock.symbol}-{r['generated']:%Y%m%d}.xlsx"
    book.save(path)
    return str(path)
